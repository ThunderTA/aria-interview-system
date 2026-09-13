"""Resume text extraction and rule-based skill / role / level inference.

Pure functions with no database or FastAPI coupling, so this module can be
unit-tested and swapped for an LLM-backed parser later without touching the
router. See docs/architecture.md ("Onboarding & Personalization").
"""

import io
import re

from app.data.skill_taxonomy import LEVEL_KEYWORDS, ROLE_SIGNALS, SKILL_TAXONOMY

MAX_RESUME_BYTES = 5 * 1024 * 1024  # 5 MB

# Regexes reused across calls.
_YEARS_PATTERN = re.compile(
    r"(\d{1,2})\+?\s*(?:\.\d+\s*)?(?:years?|yrs?)\s+(?:of\s+)?(?:experience|exp)", re.IGNORECASE
)
_WHITESPACE_PATTERN = re.compile(r"[ \t]+")
_BLANK_LINES_PATTERN = re.compile(r"\n{3,}")


class ResumeParseError(ValueError):
    """Raised when a resume file cannot be read or contains no usable text."""


def extract_text(filename: str, content: bytes) -> str:
    """Pull plain text out of a PDF or Word resume."""
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if suffix == "pdf":
        text = _extract_pdf_text(content)
    elif suffix in {"docx", "doc"}:
        text = _extract_docx_text(content)
    elif suffix == "txt":
        text = content.decode("utf-8", errors="replace")
    else:
        raise ResumeParseError(f"Unsupported file type '.{suffix}'. Upload a PDF or Word document.")

    text = _normalise_whitespace(text)
    if len(text.strip()) < 30:
        raise ResumeParseError(
            "No readable text found. If this is a scanned resume, upload a text-based PDF instead."
        )
    return text


def _extract_pdf_text(content: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except PdfReadError as exc:
        raise ResumeParseError("This PDF could not be read — it may be corrupted.") from exc


def _extract_docx_text(content: bytes) -> str:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(content))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise ResumeParseError(
            "This Word file could not be read. Older .doc files must be saved as .docx."
        ) from exc

    parts = [para.text for para in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _normalise_whitespace(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _WHITESPACE_PATTERN.sub(" ", text)
    return _BLANK_LINES_PATTERN.sub("\n\n", text).strip()


def extract_skills(text: str) -> list[dict[str, str]]:
    """Match resume text against the skill taxonomy.

    Matching is word-boundary based so short aliases ("go", "c", "r") don't fire
    inside longer words. Each canonical skill is reported at most once.
    """
    lowered = text.lower()
    found: list[dict[str, str]] = []

    for category, skills in SKILL_TAXONOMY.items():
        for canonical, aliases in skills.items():
            if any(_contains_alias(lowered, alias) for alias in aliases):
                found.append({"name": canonical, "category": category})

    return found


def _contains_alias(lowered_text: str, alias: str) -> bool:
    return re.search(rf"(?<![\w+#]){re.escape(alias)}(?![\w+#])", lowered_text) is not None


def infer_role(skills: list[dict[str, str]]) -> str:
    """Pick the interview role whose signals dominate the resume.

    Each matched skill casts a weighted vote for every role that lists its
    category as a signal (see ROLE_SIGNALS) — a rarer, more role-specific
    category (product, qa, mlops, data_analytics, hr_domain) outweighs a
    broad one (language, cs_fundamentals) that shows up on most resumes
    regardless of role. Falls back to SDE when nothing matches.

    Returns a value matching SessionRole in app/models/session.py.
    """
    scores = {role: 0.0 for role in ROLE_SIGNALS}
    for skill in skills:
        for role, weights in ROLE_SIGNALS.items():
            weight = weights.get(skill["category"])
            if weight:
                scores[role] += weight

    best_role, best_score = max(scores.items(), key=lambda kv: kv[1])
    return best_role if best_score > 0 else "SDE"


def infer_level(text: str) -> str:
    """Estimate seniority as one of: intern, junior, mid, senior."""
    lowered = text.lower()

    if any(kw in lowered for kw in LEVEL_KEYWORDS["senior"]):
        return "senior"

    years = _max_years_of_experience(lowered)
    if years is not None:
        if years >= 6:
            return "senior"
        if years >= 3:
            return "mid"
        if years >= 1:
            return "junior"

    if any(kw in lowered for kw in LEVEL_KEYWORDS["intern"]):
        return "intern"

    return "junior" if years is not None else "intern"


def _max_years_of_experience(lowered_text: str) -> int | None:
    matches = [int(m) for m in _YEARS_PATTERN.findall(lowered_text)]
    plausible = [y for y in matches if y <= 50]
    return max(plausible) if plausible else None


def parse_resume(filename: str, content: bytes) -> dict:
    """Full pipeline: bytes in, structured resume fields out."""
    text = extract_text(filename, content)
    skills = extract_skills(text)
    return {
        "raw_text": text,
        "parsed_skills": skills,
        "inferred_role": infer_role(skills),
        "inferred_level": infer_level(text),
    }
